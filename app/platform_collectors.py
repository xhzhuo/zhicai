"""Run the supplied collection engines in a per-run browser owned by the server."""
from __future__ import annotations

import asyncio
from datetime import date, datetime, timedelta
import io
import json
from pathlib import Path
import re
import tempfile
from urllib.parse import urlparse
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from playwright.async_api import async_playwright

ASSETS = Path(__file__).resolve().parent / 'tool_assets'
HEADERS = ['任务名称','启用','搜索类型','关键词','笔记类型','发布时间','开始日期','结束日期',
           '曝光量','阅读量','互动量','作者粉丝数','排序','最多笔记数']


def platform_url(value):
    value = str(value or 'https://ad.xiaohongshu.com/').strip()
    parsed = urlparse(value)
    if (parsed.scheme != 'https' or parsed.hostname != 'ad.xiaohongshu.com'
            or parsed.username or parsed.password or parsed.port not in (None, 443)):
        raise ValueError('页面链接须为 https://ad.xiaohongshu.com/ 下的聚光页面')
    return value


def read_tasks(content: bytes):
    if not content or len(content) > 20 * 1024 * 1024:
        raise ValueError('请选择不超过 20 MB 的 .xlsx 筛选条件文件')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 60 * 1024 * 1024:
                raise ValueError('文件解压后过大，请拆分任务')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('无法读取 Excel，请使用配套 .xlsx 模板') from exc
    try:
        for ws in wb:
            if ws.title != '筛选条件':
                continue
            rows = ws.iter_rows(values_only=True)
            header = [str(x or '').strip() for x in next(rows, ())]
            if not set(HEADERS).issubset(header):
                continue
            if len(set(header)) != len(header) or any(x and x not in HEADERS for x in header):
                raise ValueError('筛选条件表头重复或包含不支持的列，请使用配套模板')
            tasks = []
            for number, row in enumerate(rows, 2):
                if number > 1002:
                    raise ValueError('筛选条件表最多读取 1000 行，请删除多余空白行')
                if not any(x is not None and x != '' for x in row):
                    continue
                task = {key: row[header.index(key)] if header.index(key) < len(row) else None for key in HEADERS}
                if task['启用'] in ('否', '0', 0, False):
                    continue
                if task['启用'] not in ('是', '1', 1, True):
                    raise ValueError(f'第 {number} 行：启用须填写是或否')
                task['任务名称'] = str(task['任务名称'] or f'任务{number-1}').strip()
                task['关键词'] = str(task['关键词'] or '').strip()
                if not task['关键词'] or len(task['关键词']) > 200:
                    raise ValueError(f'第 {number} 行：请填写不超过 200 字的关键词')
                choices = {'搜索类型':['笔记','笔记ID','作者昵称','作者ID','话题','作者'],
                           '笔记类型':['不限','图文','视频'], '排序':['高曝光','高阅读','高互动'],
                           '发布时间':['近7天','近14天','近30天','近90天','近180天','近365天','自定义']}
                for key, values in choices.items():
                    if task[key] not in values:
                        raise ValueError(f'第 {number} 行：{key}填写无效')
                for key in ('开始日期','结束日期'):
                    value = task[key]
                    if isinstance(value, (datetime, date)):
                        value = value.strftime('%Y-%m-%d')
                    elif isinstance(value, (int, float)):
                        value = (date(1899,12,30) + timedelta(days=round(value))).isoformat()
                    value = str(value or '').strip()
                    if value:
                        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
                            raise ValueError(f'第 {number} 行：日期须为 YYYY-MM-DD')
                        date.fromisoformat(value)
                    task[key] = value
                if task['发布时间'] == '自定义':
                    if not task['开始日期'] or not task['结束日期'] or task['开始日期'] > task['结束日期']:
                        raise ValueError(f'第 {number} 行：请填写正确的自定义起止日期')
                elif task['开始日期'] or task['结束日期']:
                    raise ValueError(f'第 {number} 行：填写日期时须选择自定义')
                for key in ('曝光量','阅读量','互动量','作者粉丝数'):
                    task[key] = str(task[key] or '不限').strip()
                raw_limit = task['最多笔记数']
                try:
                    limit = float(1000 if raw_limit is None or raw_limit == '' else raw_limit)
                    if not limit.is_integer() or not 1 <= limit <= 10000:
                        raise ValueError()
                except (ValueError, TypeError, OverflowError):
                    raise ValueError(f'第 {number} 行：最多笔记数须为 1—10000') from None
                task['最多笔记数'] = int(limit)
                tasks.append(task)
                if len(tasks) > 100:
                    raise ValueError('一次最多执行 100 个启用任务')
            if not tasks:
                raise ValueError('没有启用的任务')
            return tasks
        raise ValueError('需要名为“筛选条件”且包含完整筛选字段的工作表，请下载配套模板')
    finally:
        wb.close()


def import_result(run, output, max_bytes):
    if output.stat().st_size > max_bytes:
        output.unlink()
        raise RuntimeError('结果文件超过服务器输出上限，请缩小采集范围')
    wb = load_workbook(output, read_only=True, data_only=True)
    try:
        summaries = []
        failures = []
        for index, ws in enumerate(wb):
            rows = ws.iter_rows(values_only=True)
            columns = [str(x or '') for x in next(rows, ())]
            preview, count = [], 0
            for row in rows:
                count += 1
                if len(preview) < 30:
                    preview.append(list(row))
                if ws.title == '采集记录' and len(row) >= 4 and row[2] != '完整':
                    failures.append({'item':str(row[1]),'error':str(row[3] or row[2])})
                if ws.title == '评论采集记录':
                    state_index = next((i for i,c in enumerate(columns) if c in ('状态','评论采集状态')), None)
                    if state_index is not None and row[state_index] != '完整':
                        failures.append({'item':'评论 / '+str(row[0]),'error':str(row[state_index])})
            summaries.append({'name':ws.title,'columns':columns,'rows':preview,'count':count})
            if index == 0:
                run.update(columns=columns, rows=preview, rowCount=count)
        run['sheets'] = summaries
        run['outputFile'] = str(output)
        run['failures'].extend(failures)
        run['status'] = 'partial' if run['failures'] else 'complete'
    finally:
        wb.close()


async def extension_browser(cookies, tool, profile):
    # Full Chromium supports extension service workers in headless mode.
    playwright = await async_playwright().start()
    path = str(ASSETS / tool)
    try:
        context = await playwright.chromium.launch_persistent_context(
            profile, channel='chromium', headless=True, accept_downloads=True,
            viewport={'width':1440,'height':1000},
            args=['--disable-extensions-except='+path, '--load-extension='+path])
        await context.add_cookies(cookies)
        worker = context.service_workers[0] if context.service_workers else await context.wait_for_event('serviceworker', timeout=15000)
        return playwright, context, worker
    except BaseException:
        await playwright.stop()
        raise


async def collect(run, cookies, folder, save, max_bytes):
    tool = run['tool']
    output = folder / '结果.xlsx'
    with tempfile.TemporaryDirectory(prefix='zhicai-'+tool+'-') as profile:
        playwright, context, worker = await extension_browser(cookies, tool, profile)
        try:
            page = context.pages[0] if context.pages else await context.new_page()
            await page.goto(run['params']['pageUrl'], wait_until='domcontentloaded', timeout=45000)
            if tool == 'adgroups':
                await collect_adgroups(run, page, worker, output, save)
            else:
                await collect_plaza(run, page, folder, output, save)
            if output.is_file():
                import_result(run, output, max_bytes)
            run['processed'] = run['total']
        finally:
            await context.close()
            await playwright.stop()


def adgroups_mode(raw):
    mode = raw.get('mode', 'all')
    if mode not in ('inspect', 'current', 'all', 'diagnostics'):
        raise ValueError('广告组操作无效')
    return mode


async def adgroup_message(worker, tab_id, command):
    return await worker.evaluate("""async ({id,type}) => chrome.tabs.sendMessage(id,{type})""",
                                 {'id':tab_id,'type':command})


def save_adgroup_state(run, state, page_total, current_count, save):
    run['adgroupState'] = {'websiteTotal':page_total, 'currentPageCount':current_count,
        'collectedCount':len(state.get('records', [])),
        'allCollected':bool(state.get('completed')) and len(state.get('records', [])) == page_total}
    run['stage'] = state.get('status', '正在读取广告组')
    save(run)


async def collect_adgroups(run, page, worker, output, save):
    mode=adgroups_mode(run.get('params', {}))
    run['stage'] = '正在识别广告组页面'; save(run)
    if mode != 'diagnostics':
        if not await page.locator('th').filter(has_text='广告组').count():
            for label in ('竞价推广','标准投','广告组'):
                candidate = page.get_by_text(label, exact=True)
                if await candidate.count():
                    await candidate.first.click(timeout=10000)
                    await asyncio.sleep(1)
        try:
            await page.locator('th').filter(has_text='广告组').first.wait_for(timeout=20000)
        except Exception as exc:
            raise RuntimeError('当前页面未找到广告组表格，无法读取可采集条数。请确认链接指向广告组列表；计划列表不适用。') from exc
    tab_id = await worker.evaluate("""async url => {
      const tabs=await chrome.tabs.query({});return tabs.find(t=>t.url===url)?.id;
    }""", page.url)
    if tab_id is None:
        raise RuntimeError('未找到聚光页面，请检查平台页面链接')
    if mode == 'diagnostics':
        response=await adgroup_message(worker,tab_id,'GET_DIAGNOSTICS')
        if not response.get('ok'):
            raise RuntimeError(response.get('error') or '无法读取分页诊断信息')
        run.update(diagnostics=response['diagnostics'],status='complete',stage='分页诊断信息已生成')
        return
    page_total = await page.evaluate("""() => {
        const match=document.body.innerText.replace(/\\s+/g,'').match(/合计(\\d+)条记录/);
        return match?Number(match[1]):null;
    }""")
    # Start from DOM rows, excluding any records passively captured from other pages.
    await adgroup_message(worker,tab_id,'CLEAR')
    current=await adgroup_message(worker,tab_id,'SCRAPE_CURRENT')
    if not current.get('ok'):
        raise RuntimeError(current.get('error') or '无法采集当前页')
    current_count=len(current.get('records',[]))
    save_adgroup_state(run,current,page_total,current_count,save)
    run['collectionScope']=mode
    if mode=='inspect':
        run['adgroupState']['collectedCount']=0
        run.update(status='complete',stage='页面条数已读取，尚未采集',processed=1,total=1)
        return
    if mode=='current':
        state=current
        if not state.get('records'):
            raise RuntimeError('当前页没有可导出的广告组数据')
    else:
        if not page_total:
            raise RuntimeError('无法读取广告组页面总记录数，已停止以避免导出不完整结果')
        task = asyncio.create_task(adgroup_message(worker,tab_id,'COLLECT_ALL'))
        try:
            while not task.done():
                await asyncio.sleep(1)
                state=await adgroup_message(worker,tab_id,'GET_STATE')
                run['processed']=len(state.get('records', []));run['total']=page_total
                save_adgroup_state(run,state,page_total,current_count,save)
            state=await task
        finally:
            if not task.done():
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)
        records=state.get('records', [])
        save_adgroup_state(run,state,page_total,current_count,save)
        if not state.get('ok') or not state.get('completed') or len(records)!=state.get('expectedTotal') or len(records)!=page_total:
            raise RuntimeError(state.get('error') or '广告组数量校验未通过，已阻止导出全部结果')
    records=state.get('records', [])
    save_adgroup_state(run,state,page_total,current_count,save)
    # Use the supplied writer so IDs, status, account and time retain the source contract.
    writer = (ASSETS/'adgroups/xlsx.js').read_text(encoding='utf-8')
    await page.evaluate(writer)
    content = await page.evaluate("""({records,context}) => Array.from(XlsxWriter.buildWorkbook(
        records.sort((a,b)=>a.campaignGroupName.localeCompare(b.campaignGroupName,'zh-CN')),
        {...context,exportedAt:new Intl.DateTimeFormat('sv-SE',{timeZone:'Asia/Shanghai',
          year:'numeric',month:'2-digit',day:'2-digit',hour:'2-digit',minute:'2-digit',second:'2-digit'}).format(new Date())}))
    """, {'records':records, 'context':state.get('context', {})})
    output.write_bytes(bytes(content))


async def collect_plaza(run, page, folder, output, save):
    # The extension is injected on all platform pages, including the login page.
    # Its presence alone is not proof that the cooperation plaza has loaded.
    run['stage']='正在确认合作广场页面';save(run)
    try:
        await page.locator('.ugc-content-filter-wrap').wait_for(timeout=30000)
    except Exception as exc:
        actual=urlparse(page.url)
        try:
            body=await page.locator('body').inner_text(timeout=3000)
        except Exception:
            body=''
        if '短信登录' in body and '账号登录' in body:
            raise RuntimeError('聚光浏览器处于未登录状态：已保存的 Cookie 未能恢复登录。请到“平台账号”重新导入完整有效的聚光 Cookie，再重试。') from exc
        if not actual.path.startswith('/microapp/creativity/ugc'):
            raise RuntimeError(f'未进入合作广场，页面跳转到了 {actual.path or "/"}。请检查所选账号的页面链接及浏览器登录状态。') from exc
        raise RuntimeError('合作广场筛选区未加载。请检查页面访问状态、登录状态或页面结构是否变化。') from exc
    host = page.locator('#xhs-excel-assistant')
    await host.wait_for(timeout=25000)
    await host.locator('#file').set_input_files(str(folder/'筛选条件.xlsx'))
    await host.locator('#run:not([disabled])').wait_for(timeout=15000)
    if run['params']['collectComments']:
        await host.locator('#collectComments').check()
        await host.locator('#comment-concurrency').evaluate("el=>el.closest('details').open=true")
        await host.locator('#comment-concurrency').select_option('1')
        await host.locator('#comment-seconds').fill('90')
    await host.locator('#run').click()
    await asyncio.sleep(.2)
    last_snapshot = None
    while True:
        snapshot = await page.evaluate("""() => {
          document.dispatchEvent(new Event('zhicai-snapshot'));
          return JSON.parse(document.documentElement.getAttribute('data-zhicai-snapshot')||'{}');
        }""")
        if snapshot.get('error'):
            raise RuntimeError(snapshot['error'])
        last_snapshot = snapshot
        run['stage'] = snapshot.get('commentStatus') if snapshot.get('commentsActive') else snapshot.get('status','采集笔记')
        audit = snapshot.get('audit', [])[1:]
        run['collectionAudit']=[{'task':str(row[1]),'state':str(row[2]),'message':str(row[3] or '')}
                                for row in audit if len(row)>=4]
        run['processed'] = sum(row[2] in ('完整','部分结果','失败') for row in audit)
        run['total'] = snapshot.get('plans') or run['total']
        rows = snapshot.get('rows', [])
        if rows:
            run.update(columns=rows[0], rows=rows[1:31], rowCount=snapshot.get('rowCount', max(0,len(rows)-1)))
        save(run)
        # Original controller transitions synchronously into comments before yielding control.
        if not snapshot.get('running') and not snapshot.get('commentsActive'):
            break
        await asyncio.sleep(1)
    status=(last_snapshot or {}).get('status','')
    if not last_snapshot or len(last_snapshot.get('rows', [])) < 2:
        audit=(last_snapshot or {}).get('audit', [])[1:]
        if status.startswith('已停止：') or any(len(row)>2 and row[2]=='失败' for row in audit):
            reason=status.split('。已采集内容仍可导出')[0]
            if not reason:
                reason=next((str(row[3]) for row in audit if len(row)>3 and row[3]),'未获得笔记数据')
            raise RuntimeError(reason)
        if not audit:
            raise RuntimeError(status or '未获得采集记录，请检查聚光连接及筛选文件')
    if not last_snapshot or len(last_snapshot.get('rows', [])) < 2:
        # Empty valid results still carry the original audit sheet.
        if not last_snapshot or len(last_snapshot.get('audit', [])) < 2:
            raise RuntimeError('未获得采集记录，请检查聚光连接及筛选文件')
    if '评论补采准备失败' in last_snapshot.get('status', ''):
        run.setdefault('failures', []).append({'item':'评论','error':last_snapshot['status']})
        await host.locator('#collectComments').uncheck()
    if not await host.locator('#export').is_enabled():
        raise RuntimeError((last_snapshot or {}).get('status') or '采集尚未产生可导出结果')
    async with page.expect_download(timeout=30000) as download_info:
        await host.locator('#export').click()
    download = await download_info.value
    await download.save_as(str(output))
