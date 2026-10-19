"""Local API tests and real Chromium tests with synthetic platform responses only."""
import asyncio
import io
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
from unittest.mock import AsyncMock
from types import SimpleNamespace

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import main, auth, platform_collectors as pc


def workbook(**changes):
    row = dict(zip(pc.HEADERS, ['测试任务','是','笔记','测试词','不限','近7天','','','不限','不限','不限','不限','高曝光',3]))
    row.update(changes)
    wb=Workbook();ws=wb.active;ws.title='筛选条件';ws.append(pc.HEADERS);ws.append([row[key] for key in pc.HEADERS])
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()


class CollectorApiTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ, {'ZHICAI_DEPLOYMENT':'local'})
        self.env.start()
        self.client=TestClient(main.app, base_url='http://127.0.0.1:8765',client=('127.0.0.1',50000))

    def tearDown(self):
        self.env.stop()

    def test_urls_do_not_accept_other_hosts_or_embedded_credentials(self):
        self.assertEqual(pc.platform_url(''), 'https://ad.xiaohongshu.com/')
        for value in ('https://evil.test/', 'https://ad.xiaohongshu.com.evil.test/',
                      'https://user:secret@ad.xiaohongshu.com/', 'http://ad.xiaohongshu.com/',
                      'https://ad.xiaohongshu.com:9999/'):
            with self.assertRaises(ValueError):pc.platform_url(value)

    def test_workbook_validation(self):
        tasks=pc.read_tasks(workbook())
        self.assertEqual(tasks[0]['关键词'], '测试词')
        self.assertEqual(tasks[0]['最多笔记数'], 3)
        for changes in ({'启用':'否'}, {'最多笔记数':1.5}, {'发布时间':'自定义'},
                        {'开始日期':'2026-10-01'}, {'搜索类型':'错误'}):
            with self.assertRaises(ValueError):pc.read_tasks(workbook(**changes))

    def test_inspection_and_template_require_new_permission(self):
        with patch.object(auth,'require_tool',side_effect=HTTPException(403,'无工具权限')):
            self.assertEqual(self.client.get('/api/plaza/template').status_code,403)
            self.assertEqual(self.client.post('/api/plaza/inspect',files={'file':('x.xlsx',workbook())}).status_code,403)
        with patch.object(auth,'require_tool',return_value={'id':'test','team_id':'t'}):
            response=self.client.post('/api/plaza/inspect',files={'file':('x.xlsx',workbook())})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()['taskCount'],1)
            self.assertEqual(self.client.post('/api/plaza/inspect',files={'file':('x.txt',b'bad')}).status_code,422)
            self.assertEqual(self.client.get('/api/plaza/template').status_code,200)

    def test_new_runs_use_existing_scheduler_and_uploaded_file(self):
        with tempfile.TemporaryDirectory() as temp, patch.object(main,'RUNS',Path(temp)), \
             patch.object(main,'account_get',return_value={'id':'test-account','service':'jug','disabled':False,'alias':'测试连接'}), \
             patch.object(main,'submit_run') as submit, patch.object(main.task_control,'check'), \
             patch.object(main,'check_disk_space'), patch.object(auth,'require_tool',return_value={'id':'test','team_id':'t'}):
            response=self.client.post('/api/runs',json={'tool':'adgroups','accountId':'test-account','params':{}})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(submit.call_args[0][0]['tool'],'adgroups')
            response=self.client.post('/api/plaza/runs',data={'account_id':'test-account','collect_comments':'true'},files={'file':('x.xlsx',workbook())})
            self.assertEqual(response.status_code,200,response.text)
            run=submit.call_args[0][0]
            self.assertTrue(run['params']['collectComments'])
            self.assertEqual((Path(temp)/run['id']/'筛选条件.xlsx').read_bytes(),workbook())

    def test_cross_team_account_cannot_be_used(self):
        with patch.object(auth,'server_mode',return_value=True), patch.object(main,'account_get',return_value={'teamId':'other'}):
            with self.assertRaises(HTTPException) as caught:
                main.enqueue_platform_collection({'id':'u','team_id':'one'},'a','adgroups',{})
            self.assertEqual(caught.exception.status_code,403)

    def test_native_result_counts_preview_and_ownership(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);run_id='c'*32;folder=root/run_id;folder.mkdir()
            output=folder/'结果.xlsx'
            wb=Workbook();ws=wb.active;ws.title='data';ws.append(['笔记ID'])
            for i in range(45):ws.append([str(10000000000000000000+i)])
            audit=wb.create_sheet('采集记录');audit.append(['任务序号','任务名称','状态','说明']);audit.append([1,'任务','部分结果','达到上限'])
            wb.save(output)
            run={'id':run_id,'tool':'plaza','rows':[],'columns':[],'sheets':[],'failures':[], 'ownerUserId':'owner'}
            pc.import_result(run,output,200000)
            self.assertEqual(run['rowCount'],45);self.assertEqual(len(run['rows']),30)
            self.assertEqual(run['status'],'partial')
            with patch.object(main,'RUNS',root), patch.object(main,'read_run',return_value=run):
                response=self.client.get('/api/runs/'+run_id)
                self.assertEqual(response.json()['rowCount'],45)
                self.assertTrue(response.json()['outputReady'])
                self.assertEqual(self.client.get('/api/runs/'+run_id+'/download').status_code,200)
                with patch.object(main,'assert_local'), patch.object(auth,'server_mode',return_value=True), patch.object(auth,'user_from_request',return_value={'id':'other','role':'member'}):
                    with self.assertRaises(HTTPException) as caught:
                        main.download(run_id,object())
                    self.assertEqual(caught.exception.status_code,404)


class BrowserCollectorTests(unittest.IsolatedAsyncioTestCase):
    async def test_cancellation_closes_collection_browser(self):
        entered=asyncio.Event()
        async def pending(*args,**kwargs):
            entered.set();await asyncio.Event().wait()
        context=SimpleNamespace(pages=[SimpleNamespace(goto=pending)],close=AsyncMock())
        playwright=SimpleNamespace(stop=AsyncMock())
        with tempfile.TemporaryDirectory() as temp, patch.object(pc,'extension_browser',new=AsyncMock(return_value=(playwright,context,None))):
            task=asyncio.create_task(pc.collect({'tool':'adgroups','params':{'pageUrl':'https://ad.xiaohongshu.com/'}},[],Path(temp),lambda _:None,1000))
            await entered.wait();task.cancel()
            with self.assertRaises(asyncio.CancelledError):await task
        context.close.assert_awaited_once();playwright.stop.assert_awaited_once()

    async def test_original_adgroups_pagination_and_text_ids(self):
        with tempfile.TemporaryDirectory() as temp:
            p,context,worker=await pc.extension_browser([], 'adgroups', str(Path(temp)/'profile'))
            try:
                calls=[]
                async def route(request_route):
                    request=request_route.request
                    if '/fixture-api' in request.url:
                        page=int(json.loads(request.post_data)['page'])
                        calls.append(page)
                        record={'campaignGroupId':str(12345678901234567890000+page),'campaignGroupName':'组'+str(page),'enable':1}
                        await request_route.fulfill(json={'data':{'list':[record],'totalCount':2,'totalPage':2,'currentPage':page}})
                    else:
                        await request_route.fulfill(content_type='text/html; charset=utf-8',body='''<html><body><div style="position:absolute;right:10px;top:10px">TEST_FEEDS</div>
                        <p>合计 2 条记录</p><table><thead><tr><th>广告组</th><th>广告组ID</th><th>广告组状态</th></tr></thead><tbody></tbody></table>
                        <script>setTimeout(()=>fetch('/fixture-api',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({page:1})}),300)</script></body></html>''')
                await context.route('https://ad.xiaohongshu.com/**', route)
                page=context.pages[0]
                await page.goto('https://ad.xiaohongshu.com/fixture')
                await asyncio.sleep(1)
                run={'processed':0,'total':1}
                output=Path(temp)/'output.xlsx'
                await pc.collect_adgroups(run,page,worker,output,lambda _:None)
                wb=load_workbook(output);rows=list(wb.active.values);wb.close()
                self.assertEqual(len(rows),3)
                self.assertIn('12345678901234567890001',str(rows))
                self.assertEqual(calls,[1,1,2])
                # A changed page total must block any subsequent incomplete export.
                await page.locator('p').evaluate("e=>e.textContent='合计 3 条记录'")
                with self.assertRaises(RuntimeError):
                    await pc.collect_adgroups(run,page,worker,Path(temp)/'blocked.xlsx',lambda _:None)
                self.assertFalse((Path(temp)/'blocked.xlsx').exists())
            finally:
                await context.close();await p.stop()

    async def test_plaza_original_engine_export_and_audit(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            assets=root/'assets';shutil.copytree(pc.ASSETS,assets)
            # A deterministic site adapter replaces only third-party DOM controls.
            # The source workbook parser, capture, Collector, controller and writer are real.
            (assets/'plaza/site-adapter.js').write_text('''globalThis.XHSSiteAdapter={available:()=>true,apply:async t=>({button:document.querySelector('#search'),validate:async()=>({关键词:t.关键词})})};''')
            with patch.object(pc,'ASSETS',assets):
                p,context,worker=await pc.extension_browser([], 'plaza', str(root/'profile'))
                try:
                    async def route(request_route):
                        if '/fixture-api' in request_route.request.url:
                            await request_route.fulfill(json={'success':True,'code':0,'data':{'cooperationCodeNoteDtoList':[{'noteId':'abcdef123456789012345678','title':'测试笔记','cmtNum':0}],'pageRespDTO':{'totalCount':1,'pageIndex':1}}})
                        else:
                            await request_route.fulfill(content_type='text/html; charset=utf-8',body='''<html><body><button id="search" onclick="fetch('/fixture-api').then(r=>r.json())">搜索</button></body></html>''')
                    await context.route('https://ad.xiaohongshu.com/**',route)
                    page=context.pages[0]
                    await page.goto('https://ad.xiaohongshu.com/microapp/creativity/ugc')
                    (root/'筛选条件.xlsx').write_bytes(workbook())
                    run={'params':{'collectComments':True},'total':1}
                    output=root/'output.xlsx'
                    await pc.collect_plaza(run,page,root,output,lambda _:None)
                    result={'failures':[]}
                    pc.import_result(result,output,200000)
                    self.assertEqual(result['status'],'complete')
                    self.assertEqual(result['rowCount'],1)
                    self.assertEqual([s['name'] for s in result['sheets']],['data','评论明细','评论词云','评论采集记录','采集记录'])
                    self.assertIn('abcdef123456789012345678',str(result['rows']))
                finally:
                    await context.close();await p.stop()


if __name__=='__main__':unittest.main()
