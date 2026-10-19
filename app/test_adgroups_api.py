import asyncio
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
from openpyxl import load_workbook
from . import adgroups_api as ag, main


def response(page, total=48):
    return {'code':0,'success':True,'data':{'totalCount':total,'pageNum':page,
        'dataList':[{'campaignGroupId':12345678901234567890+i,'campaignGroupName':'=组'+str(i),'enable':i%2}
                    for i in range((page-1)*20,min(page*20,total))]}}


class DirectAdgroupTests(unittest.IsolatedAsyncioTestCase):
    async def test_all_pages_fixed_headers_text_ids_and_formula_names(self):
        async def post(path, body):
            self.assertEqual(path,ag.PATH)
            self.assertEqual(body['startTime'],'2026-10-01')
            self.assertEqual(body['endTime'],'2026-10-19')
            return response(body['page']['pageIndex'])
        with tempfile.TemporaryDirectory() as temp:
            output=Path(temp)/'out.xlsx'
            run={'params':{'start':'2026-10-01','end':'2026-10-19'},'accountAlias':'账号','failures':[]}
            await ag.collect(run,post,output,lambda _:None,1000000)
            self.assertEqual(run['rowCount'],48)
            self.assertEqual(run['total'],3)
            self.assertEqual(len(run['rows']),30)
            wb=load_workbook(output)
            ws=wb.active
            self.assertEqual([c.value for c in ws[1]],['账号','广告组','广告组ID','广告组状态','导出时间'])
            self.assertEqual((ws.max_row,ws.max_column),(49,5))
            self.assertEqual(ws['C2'].value,'12345678901234567890')
            self.assertEqual(ws['B2'].data_type,'s')
            self.assertEqual(ws['D2'].value,'暂停')
            self.assertEqual(ws['D3'].value,'有效')
            self.assertEqual(len({r[4].value for r in list(ws.rows)[1:]}),1)
            wb.close()

    async def test_zero_results_still_export_headers(self):
        with tempfile.TemporaryDirectory() as temp:
            output=Path(temp)/'out.xlsx'
            run={'params':{},'accountAlias':'a'}
            with patch.object(ag,'fetch_page',new=AsyncMock(return_value={'totalCount':0,'dataList':[]})):
                await ag.collect(run,None,output,lambda _:None,100000)
            self.assertEqual(run['status'],'complete')
            self.assertEqual(run['rowCount'],0)
            self.assertTrue(output.is_file())

    async def test_pagination_drift_duplicates_and_empty_page_do_not_export(self):
        for second in (response(2,49),{'code':0,'data':{'pageNum':2,'totalCount':48,'dataList':[]}},response(1)):
            post=AsyncMock(side_effect=[response(1),second])
            with tempfile.TemporaryDirectory() as temp:
                output=Path(temp)/'out.xlsx'
                with self.assertRaises(RuntimeError):
                    await ag.collect({'params':{'start':'2026-10-19','end':'2026-10-19'},'accountAlias':'a'},post,output,lambda _:None,100000)
                self.assertFalse(output.exists())
        duplicated=response(2)
        duplicated['data']['dataList'][0]['campaignGroupId']=response(1)['data']['dataList'][0]['campaignGroupId']
        with self.assertRaisesRegex(RuntimeError,'重复'):
            with tempfile.TemporaryDirectory() as temp:
                await ag.collect({'params':{'start':'2026-10-19','end':'2026-10-19'},'accountAlias':'a'},AsyncMock(side_effect=[response(1),duplicated]),Path(temp)/'out.xlsx',lambda _:None,100000)

    async def test_business_errors_and_unknown_status_are_not_success(self):
        with self.assertRaisesRegex(RuntimeError,'无权限'):
            await ag.fetch_page(AsyncMock(return_value={'code':401,'msg':'无权限'}),{'start':'2026-10-19','end':'2026-10-19'},1)
        with self.assertRaises(RuntimeError):ag.status_text({'enable':9})
        for raw in ({'start':'2026-02-30','end':'2026-03-01'},{'end':'2026-10-19'},{'start':'2026-10-20','end':'2026-10-19'}):
            with self.assertRaises(ValueError):ag.validate_dates(raw)

    async def test_inspection_and_execution_never_start_browser(self):
        with patch.object(main,'platform_post',new=AsyncMock(return_value=response(1))) as post, \
             patch.object(main,'ensure_page',new=AsyncMock(side_effect=AssertionError('browser forbidden'))) as browser:
            summary=await main.read_adgroup_page_summary('a',{'start':'2026-10-19','end':'2026-10-19'})
            self.assertEqual(summary['websiteTotal'],48)
            self.assertTrue(post.await_args.kwargs['direct'])
            browser.assert_not_awaited()
        with tempfile.TemporaryDirectory() as temp,patch.object(main,'RUNS',Path(temp)), \
             patch.dict(main.locks,{'jug':asyncio.Lock()}),patch.object(main,'save_run'), \
             patch.object(main,'platform_post',new=AsyncMock(side_effect=[response(1),response(2),response(3)])), \
             patch.object(main,'ensure_page',new=AsyncMock(side_effect=AssertionError('browser forbidden'))) as browser:
            run={'id':'a','tool':'adgroups','accountId':'account','accountAlias':'账号','params':{'start':'2026-10-19','end':'2026-10-19'},'failures':[]}
            await main.execute_platform_collection(run)
            self.assertEqual(run['status'],'complete')
            self.assertEqual(run['rowCount'],48)
            browser.assert_not_awaited()
