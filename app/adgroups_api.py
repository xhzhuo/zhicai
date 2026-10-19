"""Direct, read-only Jug campaign-group API and fixed Excel output."""
from datetime import date, datetime
from zoneinfo import ZoneInfo
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

PATH = '/api/light/campaigngroup/data/list'
HEADERS = ['账号', '广告组', '广告组ID', '广告组状态', '导出时间']
PAGE_SIZE = 20


def validate_dates(raw):
    if any(key in raw for key in ('mode', 'pageUrl')):
        raise ValueError('广告组已改为接口取数，请刷新页面后提交')
    start, end = str(raw.get('start') or '').strip(), str(raw.get('end') or '').strip()
    if not start and not end:
        start = end = datetime.now(ZoneInfo('Asia/Shanghai')).date().isoformat()
    if not start or not end:
        raise ValueError('请同时选择开始日期和结束日期，或全部留空使用今天')
    try:
        for value in (start, end):
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
                raise ValueError()
            date.fromisoformat(value)
    except ValueError:
        raise ValueError('日期须为有效的 YYYY-MM-DD') from None
    if start > end:
        raise ValueError('开始日期不能晚于结束日期')
    return {'start': start, 'end': end}


async def fetch_page(post, params, index):
    result = await post(PATH, {
        'page': {'pageIndex': index, 'pageSize': PAGE_SIZE}, 'filter': {},
        'columns': ['fee', 'impression', 'click', 'ctr'],
        'startTime': params['start'], 'endTime': params['end'], 'dataCaliber': 0,
    })
    if not isinstance(result, dict) or result.get('code') != 0 or result.get('success') is False:
        message = result.get('msg') or result.get('message') if isinstance(result, dict) else None
        raise RuntimeError('广告组接口返回失败：' + str(message or '响应无效'))
    data = result.get('data')
    if not isinstance(data, dict) or not isinstance(data.get('dataList'), list):
        raise RuntimeError('广告组接口返回的数据结构无效')
    total = data.get('totalCount')
    if type(total) is not int or total < 0:
        raise RuntimeError('广告组接口未返回有效总数')
    if data.get('pageNum') != index:
        raise RuntimeError('广告组接口返回页码不符，请重试')
    return data


def status_text(record):
    for key in ('campaignGroupStatus', 'statusName', 'enableStatusName', 'stateName'):
        if isinstance(record.get(key), str) and record[key].strip():
            return record[key].strip()
    value = record.get('enable', record.get('state'))
    if value in (0, False):
        return '暂停'
    if value in (1, True):
        return '有效'
    raise RuntimeError('广告组接口返回未知状态，无法准确导出')


async def collect(run, post, output, save, max_bytes):
    records, seen = [], set()
    index, expected = 1, None
    while True:
        data = await fetch_page(post, run['params'], index)
        if expected is None:
            expected = data['totalCount']
            if expected > 100000:
                raise RuntimeError('广告组数量超过单次导出上限 100000 条')
            run['total'] = max(1, (expected + PAGE_SIZE - 1) // PAGE_SIZE)
        elif data['totalCount'] != expected:
            raise RuntimeError('采集期间广告组总数发生变化，请重试')
        items = data['dataList']
        if not items and len(records) < expected:
            raise RuntimeError('广告组分页提前结束，请重试')
        for item in items:
            if not isinstance(item, dict) or item.get('campaignGroupId') is None or not item.get('campaignGroupName'):
                raise RuntimeError('广告组记录缺少名称或 ID')
            identity = str(item['campaignGroupId'])
            if identity in seen:
                raise RuntimeError('广告组分页出现重复 ID，请重试')
            seen.add(identity)
            records.append([run['accountAlias'], item['campaignGroupName'], identity, status_text(item)])
        if len(records) > expected:
            raise RuntimeError('广告组记录数量超过接口总数，请重试')
        run.update(processed=index, stage=f'已读取 {len(records)} / {expected} 个广告组')
        save(run)
        if len(records) == expected:
            break
        if index >= run['total']:
            raise RuntimeError('广告组分页数量与接口总数不一致，请重试')
        index += 1
    exported = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m-%d %H:%M:%S')
    rows = [row + [exported] for row in records]
    wb = Workbook()
    ws = wb.active
    ws.title = '广告组信息'
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.data_type = 's'
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='FF2146')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center')
    for col, width in zip('ABCDE', [30, 48, 25, 18, 24]):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = ws.dimensions
    wb.save(output)
    wb.close()
    if output.stat().st_size > max_bytes:
        output.unlink()
        raise RuntimeError('导出文件超过服务器输出上限')
    run.update(columns=HEADERS[:], rows=rows[:30], rowCount=len(rows),
               sheets=[{'name':'广告组信息','columns':HEADERS[:],'rows':rows[:30],'count':len(rows)}],
               outputFile=str(output), status='complete', stage=f'已导出 {len(rows)} 个广告组',
               processed=run['total'], dataSource='api')
    save(run)
