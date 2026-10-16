import asyncio
import csv
import io
import json
import os
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from cryptography.fernet import Fernet

from app import main, requirements
from app.credentials import CredentialStore, CredentialStorageError


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        environment = {'ZHICAI_DEPLOYMENT': 'local', 'ZHICAI_CREDENTIAL_KEYS_FILE': '',
                       'ZHICAI_CREDENTIAL_KEYS': '' if os.name == 'nt' else Fernet.generate_key().decode('ascii')}
        self.environment = patch.dict(os.environ, environment)
        self.environment.start()
        self.original_runs = main.RUNS
        self.original_accounts_file = main.ACCOUNTS_FILE
        self.original_credentials = main.CREDENTIALS
        main.RUNS = main.DATA / 'test-runs'
        main.RUNS.mkdir(exist_ok=True)
        main.ACCOUNTS_FILE = main.RUNS / 'accounts.json'
        main.CREDENTIALS = CredentialStore(main.RUNS / 'credentials')
        main.ACCOUNTS_FILE.unlink(missing_ok=True)
        self.client = TestClient(main.app, base_url='http://127.0.0.1:8765')

    def tearDown(self):
        main.RUNS = self.original_runs
        main.ACCOUNTS_FILE.unlink(missing_ok=True)
        main.ACCOUNTS_FILE = self.original_accounts_file
        main.CREDENTIALS = self.original_credentials
        for run_id in ('a' * 32, 'b' * 32):
            (main.DATA / 'test-runs' / f'{run_id}.json').unlink(missing_ok=True)
        for account_id in ('jug', 'lingxi'):
            (main.DATA / 'test-runs' / 'credentials' / f'{account_id}.enc').unlink(missing_ok=True)
        self.environment.stop()

    def test_source_parameters_are_validated(self):
        self.assertEqual(main.validate_params('recommend', {'roots': '启赋\n启赋\n惠氏'})['roots'], ['启赋', '惠氏'])
        with self.assertRaises(ValueError):
            main.validate_params('combined', {'roots': '启赋', 'start': '2026-10-15', 'end': '2026-10-14'})
        with self.assertRaises(ValueError):
            main.validate_params('bubble', {'taxonomies': [{'name': '母婴奶粉', 'code': 'invalid code'}],
                'start': '2026-06-03', 'end': '2026-06-10', 'yindexList': ['readNum']})

    def test_public_origin_requires_server_mode_and_loopback_proxy(self):
        public_origin = 'https://119.3.124.39'
        local_request = SimpleNamespace(client=SimpleNamespace(host='127.0.0.1'),
                                        headers={'origin': public_origin})
        with patch.dict(os.environ, {'ZHICAI_DEPLOYMENT': 'server',
                                     'ZHICAI_PUBLIC_ORIGIN': public_origin}):
            main.assert_local(local_request)
            requirements.local_only(local_request)
            remote_request = SimpleNamespace(client=SimpleNamespace(host='198.51.100.10'),
                                             headers={'origin': public_origin})
            with self.assertRaises(HTTPException):
                main.assert_local(remote_request)
            wrong_origin = SimpleNamespace(client=SimpleNamespace(host='127.0.0.1'),
                                           headers={'origin': 'https://other.example'})
            with self.assertRaises(HTTPException):
                requirements.local_only(wrong_origin)
        with self.assertRaises(HTTPException):
            main.assert_local(local_request)

    def test_page_retains_v4_shell_and_loads_real_runtime(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('class="sidebar"', response.text)
        self.assertIn('id="account-drawer"', response.text)
        self.assertIn('/static/integration.js', response.text)
        self.assertNotIn('交互原型 · 模拟数据', response.text)

    def test_meta_returns_runtime_identity(self):
        response = self.client.get('/api/meta')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['localUser'])

    def test_multiple_accounts_have_distinct_ids_and_status(self):
        with patch.object(main, 'assert_local'):
            created = self.client.post('/api/accounts', json={'service': 'jug', 'alias': '品牌 A 聚光'}).json()
            self.assertEqual(created['service'], 'jug')
            self.assertNotEqual(created['id'], 'jug')
            result = self.client.patch(f"/api/accounts/{created['id']}", json={'disabled': True}).json()
            self.assertEqual(result['status'], 'disabled')
            self.assertEqual(main.account_get('jug')['status'], 'unverified')

    def test_cookie_import_accepts_service_cookie_and_rejects_other_domains(self):
        cookies, origins = main.parse_cookie_import('session=abc123; csrf=xyz', 'jug')
        self.assertEqual(len(cookies), 2)
        self.assertEqual(cookies[0]['url'], 'https://ad.xiaohongshu.com/')
        self.assertEqual(origins, [])
        with self.assertRaises(ValueError):
            main.parse_cookie_import('{"cookies":[{"name":"session","value":"abc","domain":"other.example"}]}', 'jug')

    def test_cookie_mode_creates_no_visible_browser_window(self):
        page = SimpleNamespace(url='about:blank', goto=AsyncMock())
        context = SimpleNamespace(pages=[page], new_page=AsyncMock())
        launch = AsyncMock(return_value=context)
        browser = SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=launch))
        with patch.object(main, 'playwright_instance', browser), patch.object(main, 'contexts', {}), \
             patch.object(main, 'pages', {}), patch.object(main, 'context_headless', {}), \
             patch.object(main, 'account_get', return_value={'method': 'cookie'}):
            asyncio.run(main.ensure_page('jug', headless=True, navigate=False))
        self.assertTrue(launch.call_args.kwargs['headless'])
        page.goto.assert_not_called()

    def test_verification_failure_does_not_imply_expired_cookie(self):
        with patch.object(main, 'assert_local'), patch.object(main, 'recommend', new=AsyncMock(side_effect=RuntimeError('请求过于频繁'))):
            response = self.client.post('/api/accounts/jug/verify', json={})
        self.assertEqual(response.status_code, 503)
        self.assertIn('不能据此判断 Cookie 已过期', response.json()['detail'])
        self.assertEqual(main.account_get('jug')['status'], 'unverified')

        with patch.object(main, 'assert_local'), patch.object(main, 'recommend', new=AsyncMock(side_effect=main.AuthenticationError('HTTP 401'))):
            response = self.client.post('/api/accounts/jug/verify', json={})
        self.assertEqual(response.status_code, 503)
        self.assertEqual(main.account_get('jug')['status'], 'expired')

    def test_old_unconfirmed_expired_state_is_reclassified(self):
        main.ACCOUNTS_FILE.write_text(json.dumps([{'id':'jug','service':'jug','alias':'聚光','method':'cookie','status':'expired','disabled':False}]), encoding='utf-8')
        account = main.account_get('jug')
        self.assertEqual(account['status'], 'unverified')
        self.assertEqual(account['statusReason'], 'legacy_unconfirmed')
        self.assertTrue(account['cookieConfigured'])

    def test_cookie_import_records_confirmation_without_raw_cookie(self):
        context = SimpleNamespace(add_cookies=AsyncMock())
        with patch.object(main, 'assert_local'), patch.object(main, 'ensure_page', new=AsyncMock(return_value=SimpleNamespace())), \
             patch.object(main, 'contexts', {'jug': context}):
            response = self.client.post('/api/accounts/jug/cookie', json={'content':'session=private-secret'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['cookieSavedAt'])
        self.assertTrue(main.account_get('jug')['cookieConfigured'])
        self.assertNotIn('private-secret', main.ACCOUNTS_FILE.read_text(encoding='utf-8'))
        stored = (main.CREDENTIALS.directory / 'jug.enc').read_bytes()
        self.assertNotIn(b'private-secret', stored)
        with patch.object(main, 'assert_local'):
            revealed = self.client.post('/api/accounts/jug/cookie/reveal', json={})
        self.assertEqual(revealed.json()['content'], 'session=private-secret')
        self.assertEqual(revealed.headers['cache-control'], 'no-store')

    def test_server_key_is_external_and_can_decrypt(self):
        key = Fernet.generate_key().decode('ascii')
        with patch.dict('os.environ', {'ZHICAI_DEPLOYMENT': 'server', 'ZHICAI_CREDENTIAL_KEYS': key}):
            main.CREDENTIALS.save('jug', 'session=server-value')
            self.assertEqual(main.CREDENTIALS.load('jug'), 'session=server-value')
        self.assertTrue((main.CREDENTIALS.directory / 'jug.enc').read_bytes().startswith(b'fernet1:'))

    def test_server_key_can_be_read_from_mounted_secret(self):
        key_file = main.RUNS / 'test-credential-keys'
        try:
            key_file.write_bytes(Fernet.generate_key())
            with patch.dict('os.environ', {'ZHICAI_DEPLOYMENT': 'server',
                                           'ZHICAI_CREDENTIAL_KEYS_FILE': str(key_file),
                                           'ZHICAI_CREDENTIAL_KEYS': ''}):
                main.CREDENTIALS.save('jug', 'session=file-value')
                self.assertEqual(main.CREDENTIALS.load('jug'), 'session=file-value')
                key_file.unlink()
                with self.assertRaisesRegex(CredentialStorageError, '无法读取服务器凭据密钥文件'):
                    main.CREDENTIALS.load('jug')
        finally:
            key_file.unlink(missing_ok=True)

    @unittest.skipUnless(os.name == 'nt', 'Windows DPAPI only')
    def test_local_cookie_can_be_migrated_to_server_key(self):
        main.CREDENTIALS.save('jug', 'session=migrate-value')
        self.assertTrue((main.CREDENTIALS.directory / 'jug.enc').read_bytes().startswith(b'dpapi1:'))
        key = Fernet.generate_key().decode('ascii')
        with patch.dict('os.environ', {'ZHICAI_CREDENTIAL_KEYS': key}):
            self.assertEqual(main.CREDENTIALS.migrate_to_environment_key(), 1)
            self.assertEqual(main.CREDENTIALS.load('jug'), 'session=migrate-value')
        self.assertTrue((main.CREDENTIALS.directory / 'jug.enc').read_bytes().startswith(b'fernet1:'))

    def test_old_cookie_can_be_recovered_into_encrypted_store(self):
        main.account_modify('jug', method='cookie', cookieConfigured=True)
        context = SimpleNamespace(cookies=AsyncMock(return_value=[{'name':'session','value':'legacy-value'}]))
        with patch.object(main, 'assert_local'), patch.object(main, 'ensure_page', new=AsyncMock(return_value=SimpleNamespace())), \
             patch.object(main, 'contexts', {'jug': context}):
            response = self.client.post('/api/accounts/jug/cookie/reveal', json={})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['content'], 'session=legacy-value')
        self.assertEqual(main.CREDENTIALS.load('jug'), 'session=legacy-value')

    def test_server_mode_uses_ephemeral_browser_context(self):
        page = SimpleNamespace(url='about:blank', goto=AsyncMock())
        context = SimpleNamespace(pages=[page], add_cookies=AsyncMock(), new_page=AsyncMock())
        browser = SimpleNamespace(new_context=AsyncMock(return_value=context))
        launch = AsyncMock(return_value=browser)
        persistent = AsyncMock()
        driver = SimpleNamespace(chromium=SimpleNamespace(launch=launch, launch_persistent_context=persistent))
        key = Fernet.generate_key().decode('ascii')
        with patch.dict('os.environ', {'ZHICAI_DEPLOYMENT': 'server', 'ZHICAI_CREDENTIAL_KEYS': key}), \
             patch.object(main, 'playwright_instance', driver), patch.object(main, 'contexts', {}), \
             patch.object(main, 'pages', {}), patch.object(main, 'context_headless', {}), \
             patch.object(main, 'browsers', {}), patch.object(main, 'account_get', return_value={'method':'cookie'}):
            asyncio.run(main.ensure_page('jug', headless=True, navigate=False))
        launch.assert_awaited_once()
        persistent.assert_not_awaited()

    def test_recommend_uses_real_response_shape_and_exports_csv(self):
        async def platform(service, path, body):
            self.assertEqual(service, 'jug')
            self.assertEqual(path, '/api/leona/rtb/tool/keyword/common/recommend')
            return {'code': 0, 'data': {'wordList': [{'keyword': '启赋3段', 'monthpv': 20},
                                                     {'keyword': '启赋2段', 'monthpv': 50}]}}

        run_id = 'a' * 32
        main.save_run({'id': run_id, 'tool': 'recommend', 'params': {'roots': ['启赋'], 'onlyContainRoot': True},
            'created': main.now(), 'ended': None, 'status': 'queued', 'processed': 0, 'total': 1,
            'columns': ['日期', '词根', '关键词', '月搜索指数', '词根内排名'], 'rows': [], 'sheets': [], 'failures': []})
        with patch.object(main, 'platform_post', side_effect=platform), patch.object(main.asyncio, 'sleep', new=AsyncMock()):
            asyncio.run(main.execute(run_id))
        result = main.read_run(run_id)
        self.assertEqual(result['status'], 'complete')
        self.assertEqual([row[2] for row in result['rows']], ['启赋2段', '启赋3段'])
        with patch.object(main, 'assert_local'):
            response = self.client.get(f'/api/runs/{run_id}/download')
        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(rows[1][2:], ['启赋2段', '50', '1'])

    def test_bubble_exports_separate_metric_sheets(self):
        run_id = 'b' * 32
        sheets = [
            {'name': '母婴奶粉-阅读量', 'columns': ['行业', '品牌/SPU', '阅读量'], 'rows': [['母婴奶粉', '品牌甲', 10]]},
            {'name': '母婴奶粉-阅读同比增长率', 'columns': ['行业', '品牌/SPU', '阅读同比增长率'], 'rows': [['母婴奶粉', '品牌甲', '5%']]},
        ]
        main.save_run({'id': run_id, 'tool': 'bubble', 'params': {}, 'created': main.now(), 'ended': main.now(),
            'status': 'complete', 'processed': 2, 'total': 2, 'columns': sheets[0]['columns'],
            'rows': sheets[0]['rows'] + sheets[1]['rows'], 'sheets': sheets, 'failures': []})
        preview = self.client.get(f'/api/runs/{run_id}').json()
        self.assertEqual(preview['rowCount'], 2)
        self.assertEqual(preview['rows'], sheets[0]['rows'])
        with patch.object(main, 'assert_local'):
            response = self.client.get(f'/api/runs/{run_id}/download')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertEqual(workbook.worksheets[1]['C1'].value, '阅读同比增长率')


if __name__ == '__main__':
    unittest.main()
